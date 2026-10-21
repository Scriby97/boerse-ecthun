"""Excel-Dateien: Gesamtexport (Inventar, Verkäufe, Auszahlungen, Kunden, Club) und Etiketten für DYMO Connect."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import reports
from .models import Article, ClubItem, ClubStatement, Payout, Person
from .pdf import price_text

CHF = '#,##0.00'
DATE = 'DD.MM.YYYY'


def _sheet(workbook: Workbook, title: str, header: list[str], rows, formats: dict[int, str] | None = None):
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(list(row))
    for col, fmt in (formats or {}).items():
        for (cell,) in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = fmt
    for i, title_text in enumerate(header, start=1):
        width = max([len(str(title_text))] + [len(str(c.value or "")) for (c,) in sheet.iter_rows(min_row=2, min_col=i, max_col=i)])
        sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    sheet.freeze_panes = "A2"
    return sheet


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def full_export() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    articles = Article.objects.select_related("category", "owner", "buyer").order_by("number", "number_suffix")
    _sheet(workbook, "Inventar", [
        "Nr.", "Kategorie", "Bezeichnung", "Grösse", "Besitzerin", "Preis", "Vermerk", "Anteil % (abweichend)",
        "Eingang", "Status", "Verkauft am", "Verkaufspreis", "Anteil Besitzerin", "Anteil Börse", "Käuferin",
    ], (
        (a.label, str(a.category), a.description, a.size, a.owner.name, a.price, a.price_note, a.owner_share,
         a.intake_date, a.get_status_display(), a.sold_date, a.sale_price, a.owner_amount, a.shop_amount,
         a.buyer.name if a.buyer else "")
        for a in articles
    ), {6: CHF, 9: DATE, 11: DATE, 12: CHF, 13: CHF, 14: CHF})

    positions = sorted(reports.positions(), key=lambda p: (p.date, p.sale_number or 0))
    _sheet(workbook, "Verkäufe", [
        "Datum", "Quittung", "Nr.", "Bezeichnung", "Grösse", "Art", "Besitzerin", "Preis", "Anteil Besitzerin",
        "Anteil Börse", "an Club",
    ], (
        (p.date, p.sale_number or "alte Liste", p.label, p.description, p.size, p.kind, p.owner, p.price,
         p.owner_amount, p.shop_amount, p.club_amount)
        for p in positions
    ), {1: DATE, 8: CHF, 9: CHF, 10: CHF, 11: CHF})

    _sheet(workbook, "Auszahlungen", ["Datum", "Besitzerin", "Art", "Betrag", "Notiz"], (
        (p.date, p.person.name, p.get_method_display(), p.amount, p.note)
        for p in Payout.objects.select_related("person").order_by("date", "pk")
    ), {1: DATE, 4: CHF})

    _sheet(workbook, "Kunden", ["Name", "Telefon", "E-Mail", "TWINT", "Adresse", "IBAN", "Anteil %", "Notizen"], (
        (p.name, p.phone, p.email, "ja" if p.twint else "", p.address, p.iban, p.owner_share, p.notes)
        for p in Person.objects.all()
    ))

    _sheet(workbook, "Clubartikel", ["Bezeichnung", "Grösse", "Preis Kunde", "an den Club", "Bestand"], (
        (c.name, c.size, c.customer_price, c.club_price, c.stock) for c in ClubItem.objects.all()
    ), {3: CHF, 4: CHF})

    _sheet(workbook, "Clubabrechnungen", ["per", "Total an den Club", "erstellt"], (
        (s.end_date, s.total, s.created_at.date()) for s in ClubStatement.objects.order_by("end_date")
    ), {1: DATE, 2: CHF, 3: DATE})
    return _save(workbook)


def dymo_labels(articles: list[Article]) -> bytes:
    """Tabelle für DYMO Connect («Daten importieren»): eine Zeile pro Etikette."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Etiketten"
    sheet.append(["Nummer", "Bezeichnung", "Groesse", "Preis", "Vermerk", "Boerse"])
    for a in articles:
        price, note = price_text(a)
        sheet.append([a.label, a.description, a.size, price, note, "Börse EC Thun"])
    return _save(workbook)
