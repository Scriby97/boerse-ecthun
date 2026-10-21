"""Import von Stephanies Inventarliste aus Numbers (.numbers), Excel (.xlsx) oder CSV.

1. read_rows() liest die Tabelle mit der Spalte «Inventar-Nr».
2. analyze() erkennt Artikel, Personen und Unstimmigkeiten – ohne etwas zu speichern.
3. apply() schreibt alles mit Stephanies Entscheidungen in die Datenbank (ersetzt den bisherigen Bestand).
"""

import csv
import datetime as dt
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from django.db import transaction
from django.utils import timezone

from .models import AppSettings, Article, Category, Payout, Person, Sale
from .money import split
from .normalize import (
    header_key,
    normalize_size,
    parse_contact,
    parse_price,
    person_key,
    text,
    to_date,
    to_int,
    to_money,
)

SUFFIXES = (".numbers", ".xlsx", ".csv")

# Spaltentitel der Liste → Feld
COLUMNS = {
    "datum": "intake_date",
    "kleidungsstück": "description",
    "grösse": "size",
    "größe": "size",
    "besitzer": "owner",
    "kontakt": "contact",
    "verkaufspreis": "price",
    "code": "code",
    "inventar-nr": "number",
    "verkauft": "sold_date",
    "nettoerlös börsenbesitzerin": "shop_amount",
    "nettoerlös börsenbesitzer": "shop_amount",
    "nettoerlös besitzer": "owner_amount",
    "käufer": "buyer",
    "status": "status",
}
REQUIRED = {"number": "Inventar-Nr", "description": "Kleidungsstück", "owner": "Besitzer", "code": "Code"}
UNKNOWN_OWNER = "Unbekannt"


class ImportFehler(Exception):
    """Die Datei passt nicht; die Meldung ist für Stephanie verständlich."""


def ensure_import_allowed() -> None:
    """Der Import ersetzt den ganzen Bestand – nur solange in der App noch nichts verkauft oder ausbezahlt wurde."""
    if Sale.objects.exists() or Payout.objects.exists():
        raise ImportFehler(
            "In der App gibt es schon Verkäufe oder Auszahlungen. Ein Import würde sie überschreiben "
            "und ist deshalb gesperrt."
        )


# --- 1. Lesen -------------------------------------------------------------------------------------


def read_rows(path: Path) -> list[list]:
    suffix = path.suffix.lower()
    if suffix == ".numbers":
        tables = _numbers_tables(path)
    elif suffix == ".xlsx":
        tables = _excel_tables(path)
    elif suffix == ".csv":
        tables = [_csv_rows(path)]
    else:
        raise ImportFehler("Bitte eine Numbers-, Excel- oder CSV-Datei wählen.")
    for rows in tables:
        if _find_header(rows) is not None:
            return rows
    raise ImportFehler("In der Datei gibt es keine Tabelle mit der Spalte «Inventar-Nr».")


def _numbers_tables(path: Path):
    from numbers_parser import Document

    doc = Document(str(path))
    for sheet in doc.sheets:
        for table in sheet.tables:
            yield table.rows(values_only=True)


def _excel_tables(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            yield [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _csv_rows(path: Path) -> list[list]:
    content = path.read_text(encoding="utf-8-sig")
    dialect = csv.Sniffer().sniff(content[:4096], delimiters=";,\t")
    return list(csv.reader(content.splitlines(), dialect))


def _find_header(rows) -> int | None:
    for i, row in enumerate(rows[:20]):
        if "inventar-nr" in (header_key(c) for c in row):
            return i
    return None


# --- 2. Prüfen ------------------------------------------------------------------------------------


def analyze(rows: list[list], categories: dict[int, int]) -> dict:
    """Ergebnis als JSON-fähiges dict: Artikel, Personen, Unstimmigkeiten, Zusammenfassung.

    categories: Code → Startnummer der bekannten Kategorien.
    """
    h = _find_header(rows)
    if h is None:
        raise ImportFehler("In der Datei gibt es keine Tabelle mit der Spalte «Inventar-Nr».")
    header = [header_key(c) for c in rows[h]]
    cols: dict[str, int] = {}
    for i, key in enumerate(header):
        field = COLUMNS.get(key)
        if field and field not in cols:
            cols[field] = i
    missing = [title for field, title in REQUIRED.items() if field not in cols]
    if missing:
        raise ImportFehler("Diese Spalten fehlen: " + ", ".join(missing))
    body = rows[h + 1 :]
    # Spalte ohne Titel, in der «Abgerechnet» steht
    settled_col = next(
        (
            i
            for i, key in enumerate(header)
            if not key and any(i < len(r) and text(r[i]).casefold() == "abgerechnet" for r in body)
        ),
        None,
    )

    def get(row, field):
        i = cols.get(field)
        return row[i] if i is not None and i < len(row) else None

    articles, empty_rows = [], 0
    for line, row in enumerate(body, start=h + 2):
        number = to_int(get(row, "number"))
        description = " ".join(text(get(row, "description")).split())
        if number is None and not description:
            if any(text(v) for v in row):
                empty_rows += 1
            continue
        price, price_note, price_ok = parse_price(get(row, "price"))
        sold = to_date(get(row, "sold_date"))
        settled = "abgerechnet" in text(get(row, "status")).casefold() or (
            settled_col is not None and settled_col < len(row) and text(row[settled_col]).casefold() == "abgerechnet"
        )
        owner_amount = to_money(get(row, "owner_amount"))
        contact = text(get(row, "contact"))
        donated = "nichts dafür" in f"{description} {contact}".casefold() or bool(
            sold and price and owner_amount == 0
        )
        articles.append(
            {
                "line": line,
                "number": number,
                "code": to_int(get(row, "code")),
                "description": description,
                "size": normalize_size(get(row, "size")),
                "owner": text(get(row, "owner")),
                "contact": contact,
                "price": _str(price),
                "price_note": price_note,
                "price_ok": price_ok,
                "price_raw": text(get(row, "price")),
                "intake_date": _iso(to_date(get(row, "intake_date"))),
                "sold_date": _iso(sold),
                "status": Article.Status.PAID if sold and settled else Article.Status.SOLD if sold else Article.Status.STOCK,
                "owner_amount": _str(owner_amount),
                "shop_amount": _str(to_money(get(row, "shop_amount"))),
                "buyer": text(get(row, "buyer")),
                "donated": donated,
            }
        )

    people = _people(articles)
    issues = _issues(articles, people, categories)
    return {
        "articles": articles,
        "people": people,
        "issues": issues,
        "summary": _summary(articles, people, empty_rows),
    }


def _people(articles) -> list[dict]:
    """Personen aus Besitzer- und Käuferspalte; pro Person alle verschiedenen Kontakttexte (neueste zuerst)."""
    found: dict[str, dict] = {}
    for a in articles:
        name = a["owner"] or UNKNOWN_OWNER
        p = found.setdefault(person_key(name), {"name": name, "contacts": {}, "owner": True})
        p["owner"] = True
        if a["contact"]:
            p["contacts"][a["contact"]] = max(p["contacts"].get(a["contact"], ""), a["intake_date"] or "")
    for a in articles:
        if a["buyer"]:
            found.setdefault(person_key(a["buyer"]), {"name": a["buyer"], "contacts": {}, "owner": False})
    people = []
    for key, p in found.items():
        options = sorted(p["contacts"], key=lambda c: p["contacts"][c], reverse=True)
        people.append({"key": key, "name": p["name"], "contact_options": options, "owner": p["owner"]})
    return sorted(people, key=lambda p: p["key"])


def _issues(articles, people, categories) -> list[dict]:
    issues = []
    by_number = defaultdict(list)
    for i, a in enumerate(articles):
        if a["number"] is not None:
            by_number[a["number"]].append(i)
    for number, rows in sorted(by_number.items()):
        if len(rows) > 1:
            issues.append({"id": f"dup-{number}", "type": "duplicate", "number": number, "rows": rows})
    for i, a in enumerate(articles):
        if a["number"] is None:
            issues.append({"id": f"nonum-{i}", "type": "no_number", "row": i})
        if a["code"] not in categories:
            issues.append({"id": f"cat-{i}", "type": "unknown_category", "row": i})
        elif a["number"] is not None and a["number"] // 1000 != categories[a["code"]] // 1000:
            issues.append({"id": f"range-{i}", "type": "range", "row": i})
        if not a["price_ok"]:
            issues.append({"id": f"price-{i}", "type": "price", "row": i})
        if a["sold_date"] and a["intake_date"] and a["sold_date"] < a["intake_date"]:
            issues.append({"id": f"date-{i}", "type": "sold_before_intake", "row": i})
        if not a["owner"]:
            issues.append({"id": f"owner-{i}", "type": "no_owner", "row": i})
    for n, p in enumerate(people):
        if len(p["contact_options"]) > 1:
            issues.append({"id": f"contact-{n}", "type": "contact", "person": n})
    return issues


def _summary(articles, people, empty_rows) -> dict:
    count = defaultdict(int)
    for a in articles:
        count[a["status"]] += 1
    prices = [Decimal(a["price"]) for a in articles if a["price"]]
    sold = [Decimal(a["price"]) for a in articles if a["price"] and a["sold_date"]]
    return {
        "articles": len(articles),
        "stock": count[Article.Status.STOCK],
        "sold": count[Article.Status.SOLD],
        "paid": count[Article.Status.PAID],
        "donated": sum(1 for a in articles if a["donated"]),
        "owners": sum(1 for p in people if p["owner"]),
        "buyers_only": sum(1 for p in people if not p["owner"]),
        "empty_rows": empty_rows,
        "value": str(sum(prices, Decimal("0"))),
        "sold_value": str(sum(sold, Decimal("0"))),
    }


def _iso(value: dt.date | None) -> str | None:
    return value.isoformat() if value else None


def _str(value: Decimal | None) -> str | None:
    return str(value) if value is not None else None


# --- 3. Übernehmen --------------------------------------------------------------------------------


@transaction.atomic
def apply(result: dict, decisions: dict[str, str]) -> dict:
    """Ersetzt alle Artikel und Personen durch den Inhalt der Liste. Vorher eine Sicherung anlegen!

    decisions: Issue-ID → Entscheidung (behaltene Zeile, korrigierter Preis/Datum, gewählter Kontakt).
    """
    ensure_import_allowed()
    articles = [dict(a) for a in result["articles"]]
    people = result["people"]
    renumber: set[int] = set()
    suffixes: dict[int, str] = {}
    skip: set[int] = set()
    contact_choice: dict[int, int] = {}

    for issue in result["issues"]:
        value = (decisions.get(issue["id"]) or "").strip()
        kind = issue["type"]
        if kind == "duplicate":
            if value.isdigit() and int(value) in issue["rows"]:
                # Nur dieser Artikel behält die Nummer, die anderen bekommen eine neue
                renumber.update(r for r in issue["rows"] if r != int(value))
            else:
                # Standard: alle behalten die Nummer (so sind sie beschriftet), ab dem zweiten mit «+»
                for n, row in enumerate(issue["rows"][1:], start=1):
                    suffixes[row] = "+" * n
        elif kind == "no_number":
            renumber.add(issue["row"])
        elif kind == "unknown_category":
            skip.add(issue["row"])
        elif kind == "price" and value:
            price = to_money(value)
            if price is not None:
                articles[issue["row"]]["price"] = str(price)
        elif kind == "sold_before_intake" and value:
            articles[issue["row"]]["sold_date"] = value
        elif kind == "contact" and value.isdigit():
            contact_choice[issue["person"]] = int(value)

    Article.objects.all().delete()
    Person.objects.all().delete()

    persons: dict[str, Person] = {}
    for n, p in enumerate(people):
        options = p["contact_options"]
        raw = options[contact_choice.get(n, 0)] if options else ""
        contact = parse_contact(raw)
        persons[p["key"]] = Person.objects.create(
            name=p["name"],
            phone=contact["phone"],
            email=contact["email"],
            twint=contact["twint"],
            notes=f"Kontakt aus der alten Liste: {raw}" if raw else "",
        )

    categories = {c.code: c for c in Category.objects.all()}
    settings = AppSettings.load()
    today = timezone.localdate()
    created, renumbered = 0, []

    def create(a: dict, number: int, suffix: str = ""):
        price = Decimal(a["price"]) if a["price"] is not None else None
        sold = dt.date.fromisoformat(a["sold_date"]) if a["sold_date"] else None
        owner = persons[person_key(a["owner"] or UNKNOWN_OWNER)]
        article = Article(
            number=number,
            number_suffix=suffix,
            category=categories[a["code"]],
            description=a["description"] or "(ohne Bezeichnung)",
            size=a["size"],
            owner=owner,
            price=price,
            price_note=a["price_note"],
            owner_share=Decimal("0") if a["donated"] else None,
            intake_date=dt.date.fromisoformat(a["intake_date"]) if a["intake_date"] else (sold or today),
            status=a["status"],
            buyer=persons.get(person_key(a["buyer"])) if a["buyer"] else None,
        )
        if sold:
            article.sold_date = sold
            article.sale_price = price
            if a["owner_amount"] is not None and a["shop_amount"] is not None:
                article.owner_amount = Decimal(a["owner_amount"])
                article.shop_amount = Decimal(a["shop_amount"])
            elif price is not None:
                article.owner_amount, article.shop_amount = split(price, article.effective_owner_share(settings))
        article.save()

    # Zuerst alle Artikel mit eindeutiger Nummer, dann die übrigen mit neuer Nummer
    suffixed = []
    for i, a in enumerate(articles):
        if i not in skip and i not in renumber:
            create(a, a["number"], suffixes.get(i, ""))
            created += 1
            if i in suffixes:
                suffixed.append((f"{a['number']}{suffixes[i]}", a["description"]))
    for i in sorted(renumber - skip):
        a = articles[i]
        number = categories[a["code"]].next_number()
        create(a, number)
        created += 1
        renumbered.append((a["number"], number, a["description"]))

    return {
        "articles": created,
        "persons": len(persons),
        "renumbered": renumbered,
        "suffixed": suffixed,
        "skipped": len(skip),
    }
