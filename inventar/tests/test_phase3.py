"""Clubabrechnung, Abrechnung für Besitzerinnen, Etiketten, Export und Sicherungen."""

import datetime as dt
import os
import shutil
import tempfile
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest import mock

from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from inventar import backup, exports, sales
from inventar.models import Article, Category, ClubItem, ClubStatement, Payout, Person, SaleItem
from inventar.pdf import club_statement_pdf, labels_pdf, payout_pdf


class ClubStatementTests(TestCase):
    def setUp(self):
        self.jacke = ClubItem.objects.get(name="Clubjacke neu", size="XS")
        self.hose = ClubItem.objects.get(name="Hose EC Thun", size="140")
        self.jacke.stock = 5
        self.jacke.save()

    def backfill(self, item, date, buyer="", size="XS"):
        return self.client.post(reverse("club_backfill"), {
            "date": date, "club_item": item.pk, "size": size, "buyer": buyer,
        })

    def test_nachtragen_aendert_den_bestand_nicht(self):
        self.backfill(self.hose, "2026-06-26", "Käuferin Julia", "L")
        self.backfill(self.jacke, "2026-07-13", "Käuferin Diana", "150")
        self.jacke.refresh_from_db()
        self.assertEqual(self.jacke.stock, 5)
        items = list(sales.open_club_items())
        self.assertEqual([(i.description, i.size, i.club_amount) for i in items],
                         [("Hose EC Thun", "L", Decimal("79.00")), ("Clubjacke neu", "150", Decimal("92.70"))])
        self.assertEqual(items[0].sale.buyer.name, "Käuferin Julia")

    def test_stichtag(self):
        self.assertEqual(sales.default_statement_end(dt.date(2026, 10, 21)), dt.date(2026, 6, 30))
        self.assertEqual(sales.default_statement_end(dt.date(2026, 12, 31)), dt.date(2026, 12, 31))
        self.assertEqual(sales.default_statement_end(dt.date(2026, 3, 1)), dt.date(2025, 12, 31))

    def test_abrechnung_nur_bis_zum_stichtag_und_nicht_doppelt(self):
        self.backfill(self.hose, "2026-06-26")
        self.backfill(self.jacke, "2026-07-13")
        statement = sales.create_club_statement(dt.date(2026, 6, 30))
        self.assertEqual((statement.total, statement.items.count()), (Decimal("79.00"), 1))
        self.assertEqual(sales.open_club_items().count(), 1)
        second = sales.create_club_statement(dt.date(2026, 12, 31))
        self.assertEqual(second.total, Decimal("92.70"))
        with self.assertRaises(sales.SaleError):
            sales.create_club_statement(dt.date(2026, 12, 31))

        pdf = club_statement_pdf(statement)
        self.assertTrue(pdf.startswith(b"%PDF"))

        sale = statement.items.get().sale
        with self.assertRaises(sales.SaleError):
            sales.cancel_sale(sale)
        sales.delete_club_statement(statement)
        self.assertEqual(sales.open_club_items().count(), 1)

    def test_vorschau_und_erstellen_ueber_die_seite(self):
        self.backfill(self.jacke, "2026-07-13", "Käuferin Diana")
        page = self.client.get(reverse("club_items"))
        self.assertContains(page, "92.70")
        preview = self.client.get(reverse("club_statement_new"), {"end_date": "2026-12-31"})
        self.assertContains(preview, "Käuferin Diana")
        self.client.post(reverse("club_statement_new"), {"end_date": "2026-12-31"})
        statement = ClubStatement.objects.get()
        response = self.client.get(reverse("club_statement_pdf", args=[statement.pk]))
        self.assertEqual(response["Content-Type"], "application/pdf")


class DocumentTests(TestCase):
    def setUp(self):
        self.anna = Person.objects.create(name="Muster Anna")
        cat = Category.objects.get(code=10)
        self.kleid = Article.objects.create(number=10001, category=cat, description="Kürkleid blau mit sehr langer Beschreibung, Glitzer, Strass und Rock in zwei Lagen",
                                            size="8–10", owner=self.anna, price=Decimal("45"))
        self.gratis = Article.objects.create(number=10002, category=cat, description="Röckli", owner=self.anna,
                                             price=Decimal("0"), price_note="gratis")
        self.spanne = Article.objects.create(number=10002, number_suffix="+", category=cat, description="Body",
                                             owner=self.anna, price=Decimal("50"), price_note="verhandelbar bis 40.00")

    def test_abrechnung_fuer_besitzerin(self):
        sale = sales.create_sale([{"kind": "artikel", "article_id": self.kleid.pk, "description": "Kürkleid",
                                   "size": "8–10", "price": "45.00"}])
        payout = sales.record_payout(self.anna, Payout.Method.TWINT, sale=sale)
        self.assertTrue(payout_pdf(payout).startswith(b"%PDF"))
        response = self.client.get(reverse("payout_pdf", args=[payout.pk]))
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("Abrechnung-Muster-Anna", response["Content-Disposition"])

    def test_etiketten_pdf_eine_seite_pro_artikel(self):
        pdf = labels_pdf(list(Article.objects.all()))
        self.assertEqual(pdf.count(b"/Type /Page\n"), 3)

    def test_etiketten_ueber_inventarfilter(self):
        response = self.client.get(reverse("labels", args=["pdf"]), {"q": "Body"})
        self.assertEqual(response["Content-Type"], "application/pdf")
        response = self.client.get(reverse("labels", args=["xlsx"]), {"status": ""})
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Nummer", "Bezeichnung", "Groesse", "Preis", "Vermerk", "Boerse"))
        prices = {r[0]: (r[3], r[4]) for r in rows[1:]}
        self.assertEqual(prices["10002"][0], "gratis")
        self.assertEqual(prices["10002+"], ("CHF 50.00", "verhandelbar"), "unterer Preis nicht auf der Etikette")
        response = self.client.get(reverse("labels", args=["pdf"]), {"q": "gibtesnicht"})
        self.assertRedirects(response, reverse("article_list"))

    def test_gesamtexport(self):
        sales.create_sale([{"kind": "eigen", "description": "Leggins", "size": "S", "price": "12.00"}])
        workbook = load_workbook(BytesIO(exports.full_export()))
        self.assertEqual(workbook.sheetnames, ["Inventar", "Verkäufe", "Auszahlungen", "Kunden", "Clubartikel",
                                               "Clubabrechnungen"])
        self.assertEqual(workbook["Inventar"].max_row, 4)
        self.assertEqual(workbook["Verkäufe"]["D2"].value, "Leggins")
        self.assertEqual(self.client.get(reverse("export_xlsx")).status_code, 200)

    def test_seiten_laden(self):
        for name in ["club_items", "export", "settings"]:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)

    def test_einstellungen_speichern(self):
        self.client.post(reverse("settings"), {
            "owner_share_default": "65", "owner_name": "Stephanie Luginbühl", "place": "Höfen",
            "receipt_thanks": "Danke!", "club_payee": "CH00\nClub",
        })
        from inventar.models import AppSettings
        self.assertEqual(AppSettings.load().receipt_thanks, "Danke!")


class BackupTests(TransactionTestCase):
    """Ohne äussere Transaktion, wie in der echten App."""

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.cloud = self.dir / "iCloud"
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)
        patcher = mock.patch.dict(os.environ, {"BOERSE_ICLOUD_DIR": str(self.cloud)})
        patcher.start()
        self.addCleanup(patcher.stop)
        override = override_settings(BACKUP_DIR=self.dir / "Sicherungen", AUTO_BACKUP=True)
        override.enable()
        self.addCleanup(override.disable)

    def test_sicherung_lokal_und_in_icloud(self):
        target = backup.backup_database("manuell")
        self.assertTrue(target.exists())
        self.assertTrue((self.cloud / target.name).exists())
        self.assertFalse(backup.backup_is_stale())
        self.assertIsNone(backup.ensure_daily_backup(), "heute schon gesichert")

    def test_nur_die_letzten_30_bleiben(self):
        folder = self.dir / "Sicherungen"
        folder.mkdir(parents=True)
        for i in range(35):
            (folder / f"2026-01-{i % 28 + 1:02d}_{i:03d}_alt.sqlite3").write_bytes(b"")
        backup.backup_database("neu")
        self.assertEqual(len(backup.list_backups()), 30)

    def test_zurueckholen(self):
        Person.objects.create(name="Vorher")
        target = backup.backup_database("test")
        Person.objects.create(name="Nachher")
        safety = backup.restore_backup(target.name)
        self.assertEqual(list(Person.objects.values_list("name", flat=True)), ["Vorher"])
        self.assertTrue(safety.exists())
        with self.assertRaises(FileNotFoundError):
            backup.restore_backup("../gibtesnicht.sqlite3")
        self.assertEqual(SaleItem.objects.count(), 0)
