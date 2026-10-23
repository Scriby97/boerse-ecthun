"""Übernahme früherer Clubverkäufe aus CSV (erfundene Namen)."""

import datetime as dt
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from inventar import club_import, sales
from inventar.models import ClubItem, ClubStatement, SaleItem

CSV = """Datum;Artikel;Grösse;Käuferin;Preis;an Club;abgerechnet
21.12.2025;Jacke Frauen Jiv;150;;95.00;95.00;ja
28.02.2026;Jacke Mondor;10-12;;55.00;55.00;ja
26.06.2026;Hose EC Thun;L;Beispiel Julia;79.00;79.00;nein
13.07.2026;Clubjacke neu;150;Beispiel Diana;95.00;92.70;nein
"""


class ClubImportTests(TestCase):
    def test_uebernahme_mit_abgerechneten_und_offenen(self):
        stock_before = {(c.name, c.size): c.stock for c in ClubItem.objects.all()}
        stats = club_import.import_rows(club_import.parse(CSV))
        self.assertEqual(stats, {"created": 4, "skipped": 0, "settled": 2})

        old = ClubStatement.objects.get()
        self.assertEqual((old.end_date, old.total), (dt.date(2026, 2, 28), Decimal("150.00")))
        open_items = list(sales.open_club_items(dt.date(2026, 12, 31)))
        self.assertEqual([(i.description, i.size, i.club_amount, i.sale.buyer.name) for i in open_items], [
            ("Hose EC Thun", "L", Decimal("79.00"), "Beispiel Julia"),
            ("Clubjacke neu", "150", Decimal("92.70"), "Beispiel Diana"),
        ])
        # Bestand bleibt; ausverkaufte Grösse wird ohne Bestand angelegt
        for item in ClubItem.objects.all():
            self.assertEqual(item.stock, stock_before.get((item.name, item.size), 0))
        self.assertTrue(ClubItem.objects.filter(name="Jacke Mondor", size="10-12", stock=0).exists())

    def test_doppelt_hochladen_wird_uebersprungen(self):
        club_import.import_rows(club_import.parse(CSV))
        stats = club_import.import_rows(club_import.parse(CSV))
        self.assertEqual((stats["created"], stats["skipped"]), (0, 4))
        self.assertEqual(SaleItem.objects.filter(kind="club").count(), 4)

    def test_fehlerhafte_datei(self):
        with self.assertRaises(club_import.ClubImportFehler):
            club_import.parse("Name;Preis\nx;1\n")
        with self.assertRaises(club_import.ClubImportFehler):
            club_import.parse("Datum;Artikel;an Club\nirgendwann;Jacke;5\n")

    def test_ueber_die_seite(self):
        response = self.client.post(reverse("club_import_csv"),
                                    {"file": SimpleUploadedFile("club.csv", CSV.encode("utf-8"))}, follow=True)
        self.assertContains(response, "4 Clubverkäufe übernommen, davon 2 schon beglichen")
        self.assertContains(response, "Beispiel Diana")
        self.assertContains(response, "Verkaufte Clubartikel (4)")


class ImportFileWithClubSheetsTests(TestCase):
    """Eine Import-Datei mit drei Blättern: Liste, Clubbestand, Clubverkäufe."""

    def write_file(self):
        import tempfile
        from pathlib import Path

        from openpyxl import Workbook

        from .test_import import sample_rows

        workbook = Workbook()
        liste = workbook.active
        liste.title = "Komplette Liste"
        for row in sample_rows():
            liste.append(row)
        bestand = workbook.create_sheet("Clubbestand")
        for row in [["Artikel", "Grösse", "Preis Kunde", "an Club", "Bestand"],
                    ["Clubjacke neu", "XS", 95, 92.7, 6], ["Jacke Mondor", "L", 65, 65, 1]]:
            bestand.append(row)
        verkaeufe = workbook.create_sheet("Clubverkäufe")
        for row in [line.split(";") for line in CSV.strip().splitlines()]:
            verkaeufe.append(row)
        path = Path(tempfile.mkdtemp()) / "Basis.xlsx"
        workbook.save(path)
        return path

    def run_import(self):
        from inventar import importer
        from inventar.models import Category

        path = self.write_file()
        result = importer.analyze(importer.read_rows(path), {c.code: c.start_number for c in Category.objects.all()})
        result["club"] = importer.read_club(path)
        return importer.apply(result, {})

    def test_alles_aus_einer_datei(self):
        stats = self.run_import()
        self.assertEqual(stats["articles"], 6)
        self.assertEqual(stats["club"], {"stock": 2, "sales": 4})
        self.assertEqual(ClubItem.objects.get(name="Clubjacke neu", size="XS").stock, 6)
        self.assertFalse(ClubItem.objects.filter(name="Jacke Männer Jiv").exists(), "alter Bestand ersetzt")
        self.assertEqual(sales.open_club_items().count(), 2)
        self.assertEqual(ClubStatement.objects.count(), 1)

    def test_erneuter_import_ersetzt_importierte_clubverkaeufe(self):
        from inventar.models import Sale

        self.run_import()
        self.run_import()
        self.assertEqual(Sale.objects.filter(source="import").count(), 4)
        self.assertEqual(ClubStatement.objects.count(), 1)
